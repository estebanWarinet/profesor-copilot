"""Extrae el contenido de material de estudio a markdown legible por el copilot.

Uso:
    python scripts/extraer_texto.py <archivo> [<salida.md>] [--formulas]

Formatos soportados: .xlsx/.xlsm (openpyxl), .xls (xlrd), .pdf (pypdf), .docx (python-docx),
.md/.txt (se copian). Si no se indica salida, imprime por consola.

En planillas, cada fila no vacia se emite con la referencia de cada celda (ej. `B12`),
para poder citar "archivo > hoja > celda" en explicaciones y correcciones.
"""
from __future__ import annotations

import argparse
import sys
from datetime import date, datetime
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")


def col_letra(idx: int) -> str:
    """0 -> A, 25 -> Z, 26 -> AA."""
    letras = ""
    idx += 1
    while idx:
        idx, resto = divmod(idx - 1, 26)
        letras = chr(65 + resto) + letras
    return letras


def formatear(valor) -> str:
    if isinstance(valor, bool):
        return "VERDADERO" if valor else "FALSO"
    if isinstance(valor, float):
        if valor.is_integer():
            return f"{int(valor):,}".replace(",", ".")
        entero, _, dec = f"{valor:,.4f}".rstrip("0").partition(".")
        return entero.replace(",", ".") + ("," + dec if dec else "")
    if isinstance(valor, int):
        return f"{valor:,}".replace(",", ".")
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")
    return " ".join(str(valor).split())


def filas_a_md(filas) -> list[str]:
    """filas: iterable de (nro_fila, [(ref_celda, valor_formateado, formula|None), ...])."""
    lineas = []
    for nro, celdas in filas:
        partes = []
        for ref, texto, formula in celdas:
            extra = f" ⟨{formula}⟩" if formula else ""
            partes.append(f"`{ref}` {texto}{extra}")
        lineas.append(f"- **F{nro}** · " + " · ".join(partes))
    return lineas


def extraer_xlsx(ruta: Path, formulas: bool) -> str:
    import openpyxl

    wb_val = openpyxl.load_workbook(ruta, data_only=True)
    wb_for = openpyxl.load_workbook(ruta, data_only=False) if formulas else None
    salida = [f"# {ruta.name}", ""]
    for ws in wb_val.worksheets:
        ws_for = wb_for[ws.title] if wb_for else None
        filas = []
        for row in ws.iter_rows():
            celdas = []
            for c in row:
                if c.value is None or (isinstance(c.value, str) and not c.value.strip()):
                    continue
                formula = None
                if ws_for is not None:
                    v = ws_for[c.coordinate].value
                    if isinstance(v, str) and v.startswith("="):
                        formula = v
                celdas.append((c.coordinate, formatear(c.value), formula))
            if celdas:
                filas.append((row[0].row, celdas))
        salida += [f"## Hoja: {ws.title}", ""]
        salida += filas_a_md(filas) or ["_(hoja vacia)_"]
        salida.append("")
    return "\n".join(salida)


def extraer_xls(ruta: Path) -> str:
    import xlrd

    wb = xlrd.open_workbook(str(ruta))
    salida = [f"# {ruta.name}", ""]
    for sh in wb.sheets():
        filas = []
        for r in range(sh.nrows):
            celdas = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                valor = cell.value
                if cell.ctype == xlrd.XL_CELL_TEXT and not valor.strip():
                    continue
                if cell.ctype == xlrd.XL_CELL_DATE:
                    valor = xlrd.xldate.xldate_as_datetime(valor, wb.datemode)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    valor = bool(valor)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    valor = "#ERROR"
                celdas.append((f"{col_letra(c)}{r + 1}", formatear(valor), None))
            if celdas:
                filas.append((r + 1, celdas))
        salida += [f"## Hoja: {sh.name}", ""]
        salida += filas_a_md(filas) or ["_(hoja vacia)_"]
        salida.append("")
    return "\n".join(salida)


def extraer_pdf(ruta: Path) -> str:
    from pypdf import PdfReader

    salida = [f"# {ruta.name}", ""]
    for i, pagina in enumerate(PdfReader(str(ruta)).pages, start=1):
        texto = (pagina.extract_text() or "").strip()
        salida += [f"## Pagina {i}", "", texto or "_(sin texto extraible: puede ser un escaneo, leer el PDF directamente)_", ""]
    return "\n".join(salida)


def extraer_docx(ruta: Path) -> str:
    import docx

    doc = docx.Document(str(ruta))
    salida = [f"# {ruta.name}", ""]
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        estilo = (p.style.name or "").lower()
        if estilo.startswith("heading") or estilo.startswith("título") or estilo.startswith("titulo"):
            nivel = "".join(ch for ch in estilo if ch.isdigit()) or "1"
            salida.append("#" * min(int(nivel) + 1, 6) + " " + p.text.strip())
        else:
            salida.append(p.text.strip())
        salida.append("")
    for t_idx, tabla in enumerate(doc.tables, start=1):
        salida += [f"## Tabla {t_idx}", ""]
        for fila in tabla.rows:
            salida.append("| " + " | ".join(" ".join(c.text.split()) for c in fila.cells) + " |")
        salida.append("")
    return "\n".join(salida)


def extraer(ruta: Path, formulas: bool = False) -> str:
    ext = ruta.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return extraer_xlsx(ruta, formulas)
    if ext == ".xls":
        return extraer_xls(ruta)
    if ext == ".pdf":
        return extraer_pdf(ruta)
    if ext == ".docx":
        return extraer_docx(ruta)
    if ext in (".md", ".txt"):
        return ruta.read_text(encoding="utf-8", errors="replace")
    raise SystemExit(f"Formato no soportado: {ext}. Para imagenes o .doc/.ppt, abrir el archivo directamente o convertirlo.")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("archivo", type=Path)
    ap.add_argument("salida", type=Path, nargs="?")
    ap.add_argument("--formulas", action="store_true", help="incluir formulas de Excel (solo .xlsx)")
    args = ap.parse_args()

    if not args.archivo.exists():
        raise SystemExit(f"No existe: {args.archivo}")
    md = extraer(args.archivo, args.formulas)
    if args.salida:
        args.salida.parent.mkdir(parents=True, exist_ok=True)
        args.salida.write_text(md, encoding="utf-8")
        print(f"OK -> {args.salida} ({len(md.splitlines())} lineas)")
    else:
        print(md)


if __name__ == "__main__":
    main()
