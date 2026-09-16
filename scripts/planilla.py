"""Manejo de planillas de examen practico (Excel).

Uso:
    python scripts/planilla.py nueva <plantilla.xls|xlsx> <carpeta_sesion> [--nombre respuesta-planilla.xlsx]
        Copia la plantilla a la sesion como .xlsx editable (convierte .xls si hace falta).

    python scripts/planilla.py leer <archivo.xlsx> [--plantilla <plantilla original>]
        Vuelca las celdas a markdown. Con --plantilla muestra solo lo que el alumno
        agrego o modifico respecto de la plantilla (ideal para corregir).
"""
from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, str(Path(__file__).parent))

from extraer_texto import col_letra, formatear  # noqa: E402


def leer_celdas(ruta: Path) -> dict[tuple[str, str], object]:
    """Devuelve {(hoja, ref): valor} con las celdas no vacias."""
    celdas: dict[tuple[str, str], object] = {}
    if ruta.suffix.lower() == ".xls":
        import xlrd

        wb = xlrd.open_workbook(str(ruta))
        for sh in wb.sheets():
            for r in range(sh.nrows):
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    if cell.ctype == xlrd.XL_CELL_TEXT and not cell.value.strip():
                        continue
                    celdas[(sh.name, f"{col_letra(c)}{r + 1}")] = cell.value
    else:
        import openpyxl

        wb = openpyxl.load_workbook(ruta, data_only=True)
        wb_f = openpyxl.load_workbook(ruta, data_only=False)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if v is None:
                        # Si el archivo nunca se abrio en Excel, las formulas no tienen valor cacheado.
                        f = wb_f[ws.title][c.coordinate].value
                        if isinstance(f, str) and f.startswith("="):
                            v = f
                    if v is None or (isinstance(v, str) and not v.strip()):
                        continue
                    celdas[(ws.title, c.coordinate)] = v
    return celdas


def xls_a_xlsx(origen: Path, destino: Path) -> None:
    import openpyxl
    import xlrd
    from openpyxl.utils import get_column_letter

    libro = xlrd.open_workbook(str(origen), formatting_info=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sh in libro.sheets():
        ws = wb.create_sheet(sh.name[:31])
        for r in range(sh.nrows):
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                valor = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    valor = xlrd.xldate.xldate_as_datetime(valor, libro.datemode)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    valor = bool(valor)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    continue
                ws.cell(row=r + 1, column=c + 1, value=valor)
        for (r0, r1, c0, c1) in sh.merged_cells:
            ws.merge_cells(start_row=r0 + 1, end_row=r1, start_column=c0 + 1, end_column=c1)
        for c, info in sh.colinfo_map.items():
            ws.column_dimensions[get_column_letter(c + 1)].width = max(info.width / 256, 4)
    wb.save(destino)


def cmd_nueva(args) -> None:
    plantilla: Path = args.plantilla
    sesion: Path = args.sesion
    if not plantilla.exists():
        raise SystemExit(f"No existe la plantilla: {plantilla}")
    sesion.mkdir(parents=True, exist_ok=True)
    destino = sesion / args.nombre
    if plantilla.suffix.lower() == ".xls":
        xls_a_xlsx(plantilla, destino)
        print(f"OK -> {destino} (convertida de .xls: se conservan textos, numeros y celdas combinadas; "
              "no el formato ni las formulas)")
    else:
        shutil.copy2(plantilla, destino)
        print(f"OK -> {destino}")


def cmd_leer(args) -> None:
    actual = leer_celdas(args.archivo)
    base = leer_celdas(args.plantilla) if args.plantilla else {}
    hojas: dict[str, list[str]] = {}
    for (hoja, ref), valor in actual.items():
        original = base.get((hoja, ref))
        if args.plantilla and original is not None and formatear(original) == formatear(valor):
            continue
        marca = "" if not args.plantilla else (" (nuevo)" if original is None else f" (antes: {formatear(original)})")
        hojas.setdefault(hoja, []).append(f"- `{ref}` {formatear(valor)}{marca}")
    titulo = "Celdas completadas/modificadas por el alumno" if args.plantilla else "Contenido de la planilla"
    print(f"# {titulo}: {args.archivo.name}\n")
    if not hojas:
        print("_(sin cambios respecto de la plantilla)_" if args.plantilla else "_(planilla vacia)_")
    for hoja, lineas in hojas.items():
        print(f"## Hoja: {hoja}\n")
        print("\n".join(lineas))
        print()


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    n = sub.add_parser("nueva", help="copiar plantilla a una sesion")
    n.add_argument("plantilla", type=Path)
    n.add_argument("sesion", type=Path)
    n.add_argument("--nombre", default="respuesta-planilla.xlsx")
    n.set_defaults(func=cmd_nueva)

    l = sub.add_parser("leer", help="volcar planilla a markdown")
    l.add_argument("archivo", type=Path)
    l.add_argument("--plantilla", type=Path)
    l.set_defaults(func=cmd_leer)

    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
